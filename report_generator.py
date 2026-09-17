"""
COREP Report Generator
- Responsive HTML report (mobile/desktop)
- Microsoft Word (.docx) report — landscape, 0.5" x 0.5" margins, big colorful charts
- Optional PDF via WeasyPrint (graceful skip if GTK3 missing)
- Strips prescription text from 'Treatment Plan' before display
- Filters placeholder drug names from drug charts
"""

import os
import re
import base64
import warnings
from datetime import datetime

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
from jinja2 import Template

warnings.filterwarnings("ignore")
sns.set_style("whitegrid")

CHART_DIR = "charts"
REPORT_DIR = "reports"
DOCX_CHART_DIR = os.path.join(REPORT_DIR, "_docx_charts")
os.makedirs(REPORT_DIR, exist_ok=True)
os.makedirs(DOCX_CHART_DIR, exist_ok=True)

from main_analysis import filter_real_drugs


# ============================================================
# Prescription stripping (for Treatment Plan display)
# ============================================================
PRESCRIPTION_PATTERNS = [
    r"\bRx\s*[:\-]?\s*[^.;\n]+",
    r"\bPrescription\s*[:\-]?\s*[^.;\n]+",
    r"\bPrescribed\s*[:\-]?\s*[^.;\n]+",
    r"\b(?:Tab|Tabs|Tablet|Tablets|Cap|Caps|Capsule|Capsules|"
    r"Syr|Syrup|Susp|Suspension|Inj|Injection|Cream|Oint|Ointment|"
    r"Drops|Spray|Inhaler|Supp|Suppository)"
    r"\.?\s+[A-Za-z][A-Za-z0-9\-]*(?:\s+[A-Za-z0-9\-]+)*"
    r"(?:\s+\d+\s*(?:mg|g|ml|mcg|IU|tab|caps?))?"
    r"(?:\s+(?:bd|tds|qid|od|nocte|prn|stat|daily|twice|thrice))?"
    r"(?:\s+(?:x|for)\s*\d+\s*(?:days?|weeks?|months?))?",
    r"\b\d+\s*(?:mg|g|ml|mcg|IU)\b[^.;\n]*",
    r"\b(?:bd|tds|qid|qds|od|nocte|prn|stat|q4h|q6h|q8h|q12h)\b",
    r"\b(?:x|for)\s*\d+\s*/\s*\d+\b",
    r"\bfor\s+\d+\s*(?:days?|weeks?|months?)\b",
    r"\bDispense[ds]?\b[^.;\n]*",
    r"\bRefer\s+to\s+(?:the\s+)?(?:Pharmacy|Pharmacist)\b[^.;\n]*",
    r"\bPharmacy\b[^.;\n]*",
    r"\bPharmacist\b[^.;\n]*",
]

DRUG_HINT_WORDS = (
    r"(?:Amoxicillin|Amoxycillin|Paracetamol|Acetaminophen|Ibuprofen|"
    r"Artemether|Lumefantrine|Artemether/Lumefantrine|ACT|Coartem|"
    r"Metronidazole|Ciprofloxacin|Azithromycin|Ceftriaxone|"
    r"ORS|Zinc|Vitamin\s*C|Folic\s*Acid|Ferrous\s*Sulphate|"
    r"Chloroquine|Quinine|Doxycycline|Albendazole|Mebendazole|"
    r"Amlodipine|Lisinopril|Losartan|Metformin|Glibenclamide|"
    r"Atorvastatin|Simvastatin|Omeprazole|Ranitidine|"
    r"Salbutamol|Prednisolone|Dexamethasone|Hydrocortisone|"
    r"Chlorpheniramine|Diphenhydramine|Promethazine|"
    r"Tetracycline|Erythromycin|Gentamicin|"
    r"Normal\s*Saline|Dextrose|Ringer)"
)
DRUG_PATTERN = re.compile(r"\b" + DRUG_HINT_WORDS + r"\b[^.;\n]*", re.IGNORECASE)
COMPILED_PATTERNS = [re.compile(p, re.IGNORECASE | re.MULTILINE)
                     for p in PRESCRIPTION_PATTERNS]


def strip_prescription(text, remove_drug_names=True):
    if pd.isna(text):
        return text
    s = str(text)
    if remove_drug_names:
        s = DRUG_PATTERN.sub("", s)
    for pat in COMPILED_PATTERNS:
        s = pat.sub("", s)
    s = re.sub(r"\s*[;\-,]\s*[;\-,]+", "; ", s)
    s = re.sub(r"\s{2,}", " ", s)
    s = re.sub(r"^[\s;,\-\.]+", "", s)
    s = re.sub(r"[\s;,\-\.]+$", "", s)
    s = re.sub(r"\b(?:and|or|then|with)\s*$", "", s, flags=re.IGNORECASE)
    return s.strip()


def clean_treatment_plans(df):
    if df is None or df.empty:
        return df
    if "Treatment Plan" not in df.columns:
        return df
    df = df.copy()
    df["Treatment Plan (Clinical)"] = df["Treatment Plan"].apply(strip_prescription)
    return df


# ============================================================
# Helpers
# ============================================================
def img_to_base64(path):
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode("utf-8")


def collect_charts():
    if not os.path.isdir(CHART_DIR):
        return []
    files = sorted(f for f in os.listdir(CHART_DIR) if f.lower().endswith(".png"))
    return [(f, os.path.join(CHART_DIR, f)) for f in files]


def pretty_name(filename):
    n = filename.rsplit(".", 1)[0]
    n = n.split("_", 1)[-1]
    return n.replace("_", " ").title()


# ============================================================
# DOCX-SPECIFIC CHART GENERATOR (BIG + COLORFUL)
# ============================================================
def _style_ax(ax, title, xlabel=None, ylabel=None):
    """Apply bold, colorful, readable styling to an axes."""
    ax.set_title(title, fontsize=18, fontweight="bold", color="#0b4a6f", pad=14)
    if xlabel: ax.set_xlabel(xlabel, fontsize=13, fontweight="bold")
    if ylabel: ax.set_ylabel(ylabel, fontsize=13, fontweight="bold")
    ax.tick_params(axis="both", labelsize=11)
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)


def _save_docx(fig, name):
    path = os.path.join(DOCX_CHART_DIR, name)
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def render_docx_charts(data):
    """Generate large, colorful versions of all charts for the Word report.
    Returns list of (caption, path)."""
    out = []
    p = data.get("patient", pd.DataFrame())
    cons = data.get("consultations", pd.DataFrame())
    nursing = data.get("nursing", pd.DataFrame())
    lab = data.get("lab_tests", pd.DataFrame())
    optical = data.get("optical", pd.DataFrame())
    pharm = data.get("pharmacy", pd.DataFrame())
    drugs = data.get("drugs", pd.DataFrame())

    # ---------------- Demographics ----------------
    if not p.empty:
        fig, axes = plt.subplots(2, 2, figsize=(16, 11))
        fig.suptitle("Patient Demographics", fontsize=22, fontweight="bold",
                     color="#0b4a6f", y=1.00)

        if "Gender" in p.columns and p["Gender"].notna().any():
            vc = p["Gender"].value_counts()
            sns.barplot(x=vc.index, y=vc.values, palette="husl",
                        ax=axes[0, 0], edgecolor="black")
            _style_ax(axes[0, 0], "Gender Distribution", ylabel="Patients")
        else:
            axes[0, 0].axis("off")

        if "Age" in p.columns and p["Age"].notna().any():
            sns.histplot(p["Age"].dropna(), bins=20, kde=True,
                         color="#e63946", ax=axes[0, 1], edgecolor="white")
            _style_ax(axes[0, 1], "Age Distribution", xlabel="Age (years)", ylabel="Count")
        else:
            axes[0, 1].axis("off")

        if "Age Group" in p.columns and p["Age Group"].notna().any():
            vc = p["Age Group"].value_counts().sort_index()
            sns.barplot(x=vc.index, y=vc.values, palette="viridis",
                        ax=axes[1, 0], edgecolor="black")
            _style_ax(axes[1, 0], "Age Group Distribution", ylabel="Patients")
        else:
            axes[1, 0].axis("off")

        if "Is Admitted" in p.columns and p["Is Admitted"].notna().any():
            vc = p["Is Admitted"].value_counts()
            axes[1, 1].pie(vc.values, labels=vc.index, autopct="%1.1f%%",
                            colors=sns.color_palette("Set2"), startangle=90,
                            textprops={"fontsize": 13, "fontweight": "bold"})
            axes[1, 1].set_title("Admission Status", fontsize=18,
                                  fontweight="bold", color="#0b4a6f")
        else:
            axes[1, 1].axis("off")

        plt.tight_layout()
        out.append(("Patient Demographics", _save_docx(fig, "01_demographics.png")))

    # ---------------- Monthly registrations ----------------
    if not p.empty and "Created At" in p.columns and p["Created At"].notna().any():
        t = p.copy()
        t["Month"] = t["Created At"].dt.to_period("M").astype(str)
        monthly = t.groupby("Month").size()
        fig, ax = plt.subplots(figsize=(16, 7))
        monthly.plot(kind="line", marker="o", color="#1d3557", linewidth=3,
                     markersize=10, markerfacecolor="#e63946",
                     markeredgecolor="white", ax=ax)
        ax.fill_between(range(len(monthly)), monthly.values, alpha=0.15, color="#1d3557")
        _style_ax(ax, "Monthly Patient Registrations",
                  xlabel="Month", ylabel="New Patients")
        plt.xticks(rotation=45)
        plt.tight_layout()
        out.append(("Monthly Patient Registrations", _save_docx(fig, "02_registrations.png")))

    # ---------------- Vitals box plots ----------------
    if not nursing.empty:
        vitals = [c for c in ["Blood Pressure Systolic", "Blood Pressure Diastolic",
                              "Pulse Rate", "Temperature",
                              "Respiratory Rate", "Oxygen Saturation"]
                  if c in nursing.columns]
        if vitals:
            fig, axes = plt.subplots(2, 3, figsize=(18, 10))
            fig.suptitle("Vital Signs Distribution", fontsize=22,
                         fontweight="bold", color="#0b4a6f")
            palette = sns.color_palette("husl", len(vitals))
            for ax, col, color in zip(axes.flatten(), vitals, palette):
                sns.boxplot(y=nursing[col].dropna(), ax=ax, color=color,
                            width=0.5, linewidth=2)
                sns.stripplot(y=nursing[col].dropna(), ax=ax, color="black",
                              size=3, alpha=0.4, jitter=True)
                _style_ax(ax, col, ylabel="Value")
            for ax in axes.flatten()[len(vitals):]:
                ax.axis("off")
            plt.tight_layout()
            out.append(("Vital Signs", _save_docx(fig, "03_vitals.png")))

    # ---------------- Top diagnoses ----------------
    if not cons.empty and "Diagnosis" in cons.columns:
        top = cons["Diagnosis"].value_counts().head(15)
        if not top.empty:
            fig, ax = plt.subplots(figsize=(16, 9))
            sns.barplot(y=top.index, x=top.values, palette="rocket",
                        ax=ax, edgecolor="black")
            _style_ax(ax, "Top 15 Diagnoses",
                      xlabel="Number of Cases", ylabel="Diagnosis")
            plt.tight_layout()
            out.append(("Top 15 Diagnoses", _save_docx(fig, "04_diagnoses.png")))

    # ---------------- Referrals ----------------
    if not cons.empty:
        refs = [c for c in ["Refer To Pharmacy", "Refer To Laboratory",
                            "Refer To Optician", "Refer To Specialist"]
                if c in cons.columns]
        if refs:
            def pos(x):
                return int(x.sum()) if x.dtype == bool else \
                       int(x.astype(str).str.lower().isin(["true", "yes", "1"]).sum())
            counts = cons[refs].apply(pos)
            fig, ax = plt.subplots(figsize=(14, 7))
            sns.barplot(x=counts.index, y=counts.values,
                        palette=["#2a9d8f", "#e9c46a", "#f4a261", "#e76f51"],
                        ax=ax, edgecolor="black")
            _style_ax(ax, "Referral Distribution", ylabel="Number of Referrals")
            plt.xticks(rotation=15)
            plt.tight_layout()
            out.append(("Referral Distribution", _save_docx(fig, "05_referrals.png")))

    # ---------------- Lab tests ----------------
    if not lab.empty:
        lab_cols = [c for c in ["Malaria Parasite", "Random Blood Sugar", "Hbsag"]
                    if c in lab.columns]
        if lab_cols:
            fig, axes = plt.subplots(1, len(lab_cols), figsize=(16, 7))
            if len(lab_cols) == 1:
                axes = [axes]
            palette = sns.color_palette("Set1", len(lab_cols))
            for ax, col, color in zip(axes, lab_cols, palette):
                vc = lab[col].value_counts(dropna=False)
                sns.barplot(x=vc.index.astype(str), y=vc.values,
                            ax=ax, color=color, edgecolor="black")
                _style_ax(ax, col, ylabel="Count")
                plt.setp(ax.get_xticklabels(), rotation=30, ha="right")
            plt.tight_layout()
            out.append(("Lab Test Results", _save_docx(fig, "06_lab.png")))

    # ---------------- Optical ----------------
    if not optical.empty and "Is Walk In" in optical.columns:
        vc = optical["Is Walk In"].value_counts()
        fig, ax = plt.subplots(figsize=(12, 8))
        ax.pie(vc.values, labels=vc.index, autopct="%1.1f%%",
                colors=sns.color_palette("Paired"),
                startangle=90, textprops={"fontsize": 14, "fontweight": "bold"})
        ax.set_title("Walk-in vs Referred (Optical)",
                     fontsize=20, fontweight="bold", color="#0b4a6f")
        plt.tight_layout()
        out.append(("Optical Walk-in vs Referred", _save_docx(fig, "07_optical.png")))

    # ---------------- Top drugs (filtered) ----------------
    if not pharm.empty and "Drug Name" in pharm.columns:
        names = filter_real_drugs(pharm["Drug Name"])
        top = names.value_counts().head(15)
        if not top.empty:
            fig, ax = plt.subplots(figsize=(16, 9))
            sns.barplot(y=top.index, x=top.values, palette="YlOrRd",
                        ax=ax, edgecolor="black")
            _style_ax(ax, "Top 15 Dispensed Drugs",
                      xlabel="Number of Dispensations", ylabel="Drug")
            plt.tight_layout()
            out.append(("Top 15 Dispensed Drugs", _save_docx(fig, "08_drugs.png")))

    # ---------------- Dispense rate ----------------
    if not pharm.empty and "Dispensed" in pharm.columns:
        vc = pharm["Dispensed"].value_counts()
        fig, ax = plt.subplots(figsize=(12, 8))
        ax.pie(vc.values, labels=vc.index, autopct="%1.1f%%",
                colors=sns.color_palette("Set3"),
                startangle=90, textprops={"fontsize": 14, "fontweight": "bold"})
        ax.set_title("Pharmacy Dispense Rate",
                     fontsize=20, fontweight="bold", color="#0b4a6f")
        plt.tight_layout()
        out.append(("Pharmacy Dispense Rate", _save_docx(fig, "09_dispense.png")))

    # ---------------- Stock by category ----------------
    if (not drugs.empty
            and {"Quantity", "Reorder Level", "Category"}.issubset(drugs.columns)):
        d = drugs.copy()
        d["Stock Status"] = np.where(d["Quantity"] <= d["Reorder Level"],
                                      "Reorder", "OK")
        fig, ax = plt.subplots(figsize=(16, 8))
        sns.barplot(data=d, x="Category", y="Quantity", hue="Stock Status",
                    palette={"OK": "#2a9d8f", "Reorder": "#e63946"},
                    ax=ax, edgecolor="black")
        _style_ax(ax, "Drug Stock by Category",
                  ylabel="Quantity on Hand", xlabel="Category")
        plt.xticks(rotation=20)
        plt.tight_layout()
        out.append(("Drug Stock by Category", _save_docx(fig, "10_stock.png")))

    return out


# ============================================================
# Summary tables
# ============================================================
def build_summary(data):
    s = {}
    p = data.get("patient", pd.DataFrame())
    c = data.get("consultations", pd.DataFrame())
    lab = data.get("lab_tests", pd.DataFrame())
    ph = data.get("pharmacy", pd.DataFrame())
    d = data.get("drugs", pd.DataFrame())

    if not p.empty:
        s["Total Patients"] = p["Id"].nunique() if "Id" in p.columns else len(p)
        if "Gender" in p.columns:
            s["Gender Split"] = ", ".join(
                f"{k}: {v}" for k, v in p["Gender"].value_counts().items())
        if "Age" in p.columns and p["Age"].notna().any():
            s["Average Age"] = f"{p['Age'].mean():.1f} years"
        if "Is Admitted" in p.columns:
            adm = p["Is Admitted"].astype(str).str.lower().isin(["true", "yes", "1"]).sum()
            s["Admitted Patients"] = int(adm)

    if not c.empty and "Diagnosis" in c.columns:
        top = c["Diagnosis"].value_counts().head(5)
        s["Top 5 Diagnoses"] = " | ".join(f"{k} ({v})" for k, v in top.items())

    if not lab.empty and "Malaria Parasite" in lab.columns:
        pos = lab["Malaria Parasite"].astype(str).str.lower().isin(
            ["positive", "pos", "reactive", "+", "1", "true"]).sum()
        s["Malaria Positive Tests"] = int(pos)

    if not ph.empty and "Drug Name" in ph.columns:
        names = filter_real_drugs(ph["Drug Name"])
        top = names.value_counts().head(5)
        if not top.empty:
            s["Top 5 Dispensed Drugs"] = " | ".join(f"{k} ({v})" for k, v in top.items())

    if not d.empty and {"Quantity", "Reorder Level"}.issubset(d.columns):
        low = (d["Quantity"] <= d["Reorder Level"]).sum()
        s["Drugs Needing Reorder"] = int(low)

    return s


# ============================================================
# HTML (unchanged responsive template)
# ============================================================
HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1.0"/>
<title>COREP Analysis Report</title>
<style>
  :root { --primary: #0b4a6f; --accent: #1a7fad; --bg: #f7f9fb;
          --card: #ffffff; --text: #223; --muted: #666; }
  * { box-sizing: border-box; }
  html, body { margin: 0; padding: 0;
    font-family: 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
    background: var(--bg); color: var(--text); line-height: 1.55; }
  .container { max-width: 1100px; margin: 0 auto; padding: clamp(16px, 3vw, 40px); }
  header.report-header {
    background: linear-gradient(135deg, var(--primary) 0%, var(--accent) 100%);
    color: #fff; padding: clamp(20px, 4vw, 40px); border-radius: 12px;
    margin-bottom: clamp(16px, 3vw, 32px);
    box-shadow: 0 4px 14px rgba(0,0,0,0.08); }
  header.report-header h1 { margin: 0 0 6px; font-size: clamp(1.4rem, 4vw, 2rem); }
  header.report-header .meta { opacity: 0.85; font-size: clamp(0.8rem, 2vw, 0.95rem); }
  h2 { color: var(--primary); font-size: clamp(1.1rem, 3vw, 1.5rem);
       border-bottom: 2px solid var(--primary); padding-bottom: 6px;
       margin: clamp(24px, 4vw, 40px) 0 16px; }
  .card { background: var(--card); padding: clamp(14px, 2.5vw, 24px);
          border-radius: 10px; box-shadow: 0 1px 4px rgba(0,0,0,0.06);
          margin-bottom: 16px; }
  table.summary { width: 100%; border-collapse: collapse;
                  font-size: clamp(0.85rem, 2vw, 1rem); }
  table.summary th, table.summary td {
    padding: 10px 12px; border-bottom: 1px solid #eee;
    text-align: left; vertical-align: top; }
  table.summary th { background: var(--primary); color: #fff;
                     font-weight: 500; width: 35%; }
  .chart-grid { display: grid;
    grid-template-columns: repeat(auto-fit, minmax(320px, 1fr));
    gap: clamp(12px, 2vw, 20px); }
  .chart-card { background: var(--card); padding: clamp(10px, 2vw, 16px);
    border-radius: 10px; box-shadow: 0 1px 4px rgba(0,0,0,0.06);
    text-align: center; overflow: hidden; }
  .chart-card img { width: 100%; height: auto; display: block;
                    border-radius: 6px; }
  .chart-card .caption { color: var(--muted);
    font-size: clamp(0.75rem, 1.8vw, 0.9rem); margin-top: 8px; font-weight: 500; }
  ul.insights { padding-left: 20px; font-size: clamp(0.9rem, 2vw, 1rem); }
  ul.insights li { margin-bottom: 8px; }
  footer.report-footer { margin-top: 40px; text-align: center;
    color: var(--muted); font-size: 0.85rem; }
  @media print {
    body { background: #fff; }
    .container { max-width: 100%; padding: 0; }
    header.report-header { box-shadow: none; border-radius: 0; }
    .chart-grid { grid-template-columns: repeat(2, 1fr); }
    .chart-card { page-break-inside: avoid; box-shadow: none;
                  border: 1px solid #eee; }
    h2 { page-break-after: avoid; } }
  @media (max-width: 480px) {
    table.summary th { width: 45%; }
    .chart-grid { grid-template-columns: 1fr; } }
</style>
</head>
<body>
<div class="container">
  <header class="report-header">
    <h1>🏥 COREP Clinical Data Analysis Report</h1>
    <div class="meta">Generated on {{ generated_at }}</div>
  </header>

  <h2>1. Executive Summary</h2>
  <div class="card">
    <table class="summary">
      {% for k, v in summary.items() %}
      <tr><th>{{ k }}</th><td>{{ v }}</td></tr>
      {% endfor %}
    </table>
  </div>

  <h2>2. Charts &amp; Findings</h2>
  <div class="chart-grid">
    {% for name, b64 in charts %}
    <div class="chart-card">
      <img src="data:image/png;base64,{{ b64 }}" alt="{{ name }}"/>
      <div class="caption">{{ name.replace('_',' ').replace('.png','').title() }}</div>
    </div>
    {% endfor %}
  </div>

  <h2>3. Key Insights</h2>
  <div class="card">
    <ul class="insights">
      <li>Patient registrations show a consistent trend — monitor monthly volume.</li>
      <li>Top diagnoses drive most of the clinical workload.</li>
      <li>Referral pathway shows the strongest downstream load.</li>
      <li>Low-stock alerts must be actioned promptly.</li>
      <li>Use forecasting to plan drug &amp; staff demand.</li>
      <li>Patient clusters reveal distinct risk groups.</li>
    </ul>
  </div>

  <footer class="report-footer">
    COREP Analytics Pipeline · {{ generated_at }}
  </footer>
</div>
</body>
</html>
"""


def generate_html_report(summary, charts):
    html = Template(HTML_TEMPLATE).render(
        generated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        summary=summary, charts=charts,
    )
    out = os.path.join(REPORT_DIR, "corep_report2.html")
    with open(out, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"📄 Responsive HTML report -> {out}")
    return out, html


# ============================================================
# WORD REPORT (LANDSCAPE, 0.5" x 0.5" MARGINS, BIG CHARTS)
# ============================================================
def generate_word_report(summary, chart_items, data=None):
    """
    chart_items: list of (caption, path) — pre-rendered, big & colorful
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor, Emu
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_ORIENT
    except ImportError:
        print("⚠️  python-docx not installed. Run: pip install python-docx")
        return None

    doc = Document()

    # ---- Landscape orientation, 0.5" x 0.5" margins ----
    for section in doc.sections:
        # Landscape
        new_w, new_h = section.page_height, section.page_width
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width  = new_w
        section.page_height = new_h
        # 0.5 inch margins
        section.top_margin    = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin   = Inches(0.5)
        section.right_margin  = Inches(0.5)

    # Usable width in a landscape Letter/A4 with 0.5" margins ≈ 10 inches
    # Charts rendered at 16" x 9" scale — set display width to 9.7"
    IMG_WIDTH = Inches(9.7)

    # ---- Title ----
    title = doc.add_heading("COREP Clinical Data Analysis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in title.runs:
        r.font.color.rgb = RGBColor(11, 74, 111)

    sub = doc.add_paragraph(
        f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}   |   "
        f"Landscape · 0.5\" margins"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(11)
    sub.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()

    # ---- Executive Summary ----
    doc.add_heading("1. Executive Summary", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for k, v in summary.items():
        row = table.add_row().cells
        row[0].text = str(k)
        row[1].text = str(v)
        for para in row[0].paragraphs:
            for r in para.runs:
                r.bold = True
                r.font.size = Pt(12)
        for para in row[1].paragraphs:
            for r in para.runs:
                r.font.size = Pt(12)

    doc.add_page_break()

    # ---- Charts ----
    doc.add_heading("2. Charts & Findings", level=1)

    for i, (caption, path) in enumerate(chart_items, start=1):
        # Chart heading
        h = doc.add_heading(f"2.{i} {caption}", level=2)
        for r in h.runs:
            r.font.color.rgb = RGBColor(11, 74, 111)

        # Big centered image
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            run = p.add_run()
            run.add_picture(path, width=IMG_WIDTH)
        except Exception as e:
            doc.add_paragraph(f"[Chart could not be embedded: {e}]")

        # Page break between charts (except after last)
        if i < len(chart_items):
            doc.add_page_break()

    # ---- Treatment Plans (prescription-free) ----
    if data is not None and not data.get("consultations", pd.DataFrame()).empty:
        cons = data["consultations"]
        if "Treatment Plan (Clinical)" in cons.columns:
            doc.add_page_break()
            doc.add_heading("3. Treatment Plans (Clinical Only)", level=1)
            note = doc.add_paragraph(
                "Prescription-related text has been removed. "
                "Only clinical/non-pharmacy instructions are shown."
            )
            note.runs[0].font.size = Pt(9)
            note.runs[0].font.color.rgb = RGBColor(120, 120, 120)

            t = doc.add_table(rows=1, cols=2)
            t.style = "Light Grid Accent 1"
            hdr = t.rows[0].cells
            hdr[0].text = "Consultation ID"
            hdr[1].text = "Treatment Plan (Clinical)"
            for c in hdr:
                for para in c.paragraphs:
                    for r in para.runs:
                        r.bold = True

            sub = cons[cons["Treatment Plan (Clinical)"].astype(str).str.strip().ne("")]
            for _, row in sub.head(20).iterrows():
                cells = t.add_row().cells
                cells[0].text = str(row.get("Id", ""))
                cells[1].text = str(row["Treatment Plan (Clinical)"])

    # ---- Key Insights ----
    doc.add_page_break()
    doc.add_heading("4. Key Insights", level=1)
    insights = [
        "Patient registrations show a consistent trend — monitor monthly volume for resource planning.",
        "Top diagnoses drive most of the clinical workload — target training and drug stocking accordingly.",
        "Referral pathway indicates the strongest downstream load (Lab vs Pharmacy vs Optical).",
        "Low-stock alerts must be actioned promptly to avoid stockouts.",
        "Use the forecasting module to plan drug & staff demand for the coming month.",
        "Patient clusters reveal distinct risk groups for targeted interventions.",
    ]
    for item in insights:
        doc.add_paragraph(item, style="List Bullet")

    out = os.path.join(REPORT_DIR, "corep_report2.docx")
    doc.save(out)
    print(f"📝 Word report (landscape, 0.5\" margins) -> {out}")
    return out


# ============================================================
# PDF (optional)
# ============================================================
def generate_pdf_report(html_str):
    try:
        from weasyprint import HTML
    except Exception as e:
        print(f"ℹ️  PDF skipped (WeasyPrint unavailable): {e}")
        return None
    try:
        out = os.path.join(REPORT_DIR, "corep_report.pdf")
        HTML(string=html_str, base_url=os.getcwd()).write_pdf(
            out, presentational_hints=True)
        print(f"📕 PDF report -> {out}")
        return out
    except Exception as e:
        print(f"ℹ️  PDF skipped (WeasyPrint failed): {e}")
        print("    → Fix: install GTK3 runtime or use conda-forge.")
        return None


# ============================================================
# Main
# ============================================================
def generate_report(data=None):
    print("\n=== GENERATING REPORT ===")

    if data is None:
        import main_analysis as analysis
        data = analysis.clean_all(analysis.load_data())

    # Strip prescriptions
    if not data.get("consultations", pd.DataFrame()).empty:
        data["consultations"] = clean_treatment_plans(data["consultations"])
        print("✅ Treatment plans cleaned (prescriptions removed).")

    summary = build_summary(data)

    # HTML uses the small (fast) 110 dpi charts from ./charts/
    chart_files = collect_charts()
    charts_b64 = [(name, img_to_base64(path)) for name, path in chart_files]

    # DOCX uses big, colorful, 200 dpi charts rendered on-the-fly
    print("🎨 Rendering large colorful charts for Word report ...")
    docx_charts = render_docx_charts(data)
    print(f"   → {len(docx_charts)} charts rendered")

    html_path, html_str = generate_html_report(summary, charts_b64)
    word_path = generate_word_report(summary, docx_charts, data=data)
    pdf_path  = generate_pdf_report(html_str)

    print("\n📦 Report outputs:")
    print(f"   • HTML: {html_path}")
    if word_path: print(f"   • Word: {word_path}")
    if pdf_path:  print(f"   • PDF : {pdf_path}")

    return {"html": html_path, "word": word_path, "pdf": pdf_path}


if __name__ == "__main__":
    generate_report()

# """
# COREP Report Generator
# - Responsive HTML report (works on mobile/desktop)
# - Microsoft Word (.docx) report
# - Optional PDF via WeasyPrint (graceful skip if GTK3 missing)
# """

# import os
# import base64
# import warnings
# from datetime import datetime
# from io import BytesIO

# import pandas as pd
# from jinja2 import Template

# warnings.filterwarnings("ignore")

# CHART_DIR = "charts"
# REPORT_DIR = "reports"
# os.makedirs(REPORT_DIR, exist_ok=True)


# # ============================================================
# # Helpers
# # ============================================================
# def img_to_base64(path):
#     """Read PNG → base64 string."""
#     with open(path, "rb") as f:
#         return base64.b64encode(f.read()).decode("utf-8")


# def collect_charts():
#     """Return sorted [(filename, full_path), ...]."""
#     if not os.path.isdir(CHART_DIR):
#         return []
#     files = sorted(f for f in os.listdir(CHART_DIR) if f.lower().endswith(".png"))
#     return [(f, os.path.join(CHART_DIR, f)) for f in files]


# def pretty_name(filename):
#     """'03_vitals_boxplots.png' → 'Vitals Boxplots'."""
#     n = filename.rsplit(".", 1)[0]
#     n = n.split("_", 1)[-1]  # drop leading number
#     return n.replace("_", " ").title()


# # ============================================================
# # Summary tables
# # ============================================================
# def build_summary(data):
#     s = {}
#     p = data.get("patient", pd.DataFrame())
#     c = data.get("consultations", pd.DataFrame())
#     lab = data.get("lab_tests", pd.DataFrame())
#     ph = data.get("pharmacy", pd.DataFrame())
#     d = data.get("drugs", pd.DataFrame())

#     if not p.empty:
#         s["Total Patients"] = p["Id"].nunique() if "Id" in p.columns else len(p)
#         if "Gender" in p.columns:
#             s["Gender Split"] = ", ".join(
#                 f"{k}: {v}" for k, v in p["Gender"].value_counts().items())
#         if "Age" in p.columns and p["Age"].notna().any():
#             s["Average Age"] = f"{p['Age'].mean():.1f} years"
#         if "Is Admitted" in p.columns:
#             adm = p["Is Admitted"].astype(str).str.lower().isin(["true", "yes", "1"]).sum()
#             s["Admitted Patients"] = int(adm)

#     if not c.empty and "Diagnosis" in c.columns:
#         top = c["Diagnosis"].value_counts().head(5)
#         s["Top 5 Diagnoses"] = " | ".join(f"{k} ({v})" for k, v in top.items())

#     if not lab.empty and "Malaria Parasite" in lab.columns:
#         pos = lab["Malaria Parasite"].astype(str).str.lower().isin(
#             ["positive", "pos", "reactive", "+", "1", "true"]).sum()
#         s["Malaria Positive Tests"] = int(pos)

#     if not ph.empty and "Drug Name" in ph.columns:
#         top = ph["Drug Name"].value_counts().head(5)
#         s["Top 5 Drugs"] = " | ".join(f"{k} ({v})" for k, v in top.items())

#     if not d.empty and {"Quantity", "Reorder Level"}.issubset(d.columns):
#         low = (d["Quantity"] <= d["Reorder Level"]).sum()
#         s["Drugs Needing Reorder"] = int(low)

#     return s


# # ============================================================
# # Responsive HTML Template
# # ============================================================
# HTML_TEMPLATE = """
# <!DOCTYPE html>
# <html lang="en">
# <head>
# <meta charset="utf-8"/>
# <meta name="viewport" content="width=device-width, initial-scale=1.0"/>
# <title>COREP Analysis Report</title>
# <style>
#   :root {
#     --primary: #0b4a6f;
#     --accent: #1a7fad;
#     --bg: #f7f9fb;
#     --card: #ffffff;
#     --text: #223;
#     --muted: #666;
#   }
#   * { box-sizing: border-box; }
#   html, body {
#     margin: 0; padding: 0;
#     font-family: 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
#     background: var(--bg);
#     color: var(--text);
#     line-height: 1.55;
#     -webkit-text-size-adjust: 100%;
#   }
#   .container {
#     max-width: 1100px;
#     margin: 0 auto;
#     padding: clamp(16px, 3vw, 40px);
#   }
#   header.report-header {
#     background: linear-gradient(135deg, var(--primary) 0%, var(--accent) 100%);
#     color: #fff;
#     padding: clamp(20px, 4vw, 40px);
#     border-radius: 12px;
#     margin-bottom: clamp(16px, 3vw, 32px);
#     box-shadow: 0 4px 14px rgba(0,0,0,0.08);
#   }
#   header.report-header h1 {
#     margin: 0 0 6px;
#     font-size: clamp(1.4rem, 4vw, 2rem);
#     font-weight: 600;
#   }
#   header.report-header .meta {
#     opacity: 0.85;
#     font-size: clamp(0.8rem, 2vw, 0.95rem);
#   }
#   h2 {
#     color: var(--primary);
#     font-size: clamp(1.1rem, 3vw, 1.5rem);
#     border-bottom: 2px solid var(--primary);
#     padding-bottom: 6px;
#     margin: clamp(24px, 4vw, 40px) 0 16px;
#   }
#   .card {
#     background: var(--card);
#     padding: clamp(14px, 2.5vw, 24px);
#     border-radius: 10px;
#     box-shadow: 0 1px 4px rgba(0,0,0,0.06);
#     margin-bottom: 16px;
#   }
#   table.summary {
#     width: 100%;
#     border-collapse: collapse;
#     font-size: clamp(0.85rem, 2vw, 1rem);
#   }
#   table.summary th, table.summary td {
#     padding: 10px 12px;
#     border-bottom: 1px solid #eee;
#     text-align: left;
#     vertical-align: top;
#   }
#   table.summary th {
#     background: var(--primary);
#     color: #fff;
#     font-weight: 500;
#     width: 35%;
#   }
#   table.summary tr:last-child th,
#   table.summary tr:last-child td { border-bottom: none; }

#   .chart-grid {
#     display: grid;
#     grid-template-columns: repeat(auto-fit, minmax(320px, 1fr));
#     gap: clamp(12px, 2vw, 20px);
#   }
#   .chart-card {
#     background: var(--card);
#     padding: clamp(10px, 2vw, 16px);
#     border-radius: 10px;
#     box-shadow: 0 1px 4px rgba(0,0,0,0.06);
#     text-align: center;
#     overflow: hidden;
#   }
#   .chart-card img {
#     width: 100%;
#     height: auto;
#     max-width: 100%;
#     display: block;
#     border-radius: 6px;
#   }
#   .chart-card .caption {
#     color: var(--muted);
#     font-size: clamp(0.75rem, 1.8vw, 0.9rem);
#     margin-top: 8px;
#     font-weight: 500;
#   }
#   ul.insights {
#     padding-left: 20px;
#     font-size: clamp(0.9rem, 2vw, 1rem);
#   }
#   ul.insights li { margin-bottom: 8px; }

#   footer.report-footer {
#     margin-top: 40px;
#     text-align: center;
#     color: var(--muted);
#     font-size: 0.85rem;
#   }

#   /* ---- Print / PDF ---- */
#   @media print {
#     body { background: #fff; }
#     .container { max-width: 100%; padding: 0; }
#     header.report-header { box-shadow: none; border-radius: 0; }
#     .chart-grid { grid-template-columns: repeat(2, 1fr); }
#     .chart-card { page-break-inside: avoid; box-shadow: none; border: 1px solid #eee; }
#     h2 { page-break-after: avoid; }
#   }

#   /* ---- Small mobile ---- */
#   @media (max-width: 480px) {
#     table.summary th { width: 45%; }
#     .chart-grid { grid-template-columns: 1fr; }
#   }
# </style>
# </head>
# <body>
# <div class="container">

#   <header class="report-header">
#     <h1>🏥 COREP Clinical Data Analysis Report</h1>
#     <div class="meta">Generated on {{ generated_at }}</div>
#   </header>

#   <h2>1. Executive Summary</h2>
#   <div class="card">
#     <table class="summary">
#       {% for k, v in summary.items() %}
#       <tr><th>{{ k }}</th><td>{{ v }}</td></tr>
#       {% endfor %}
#     </table>
#   </div>

#   <h2>2. Charts &amp; Findings</h2>
#   <div class="chart-grid">
#     {% for name, b64 in charts %}
#     <div class="chart-card">
#       <img src="data:image/png;base64,{{ b64 }}" alt="{{ name }}"/>
#       <div class="caption">{{ name.replace('_',' ').replace('.png','').title() }}</div>
#     </div>
#     {% endfor %}
#   </div>

#   <h2>3. Key Insights</h2>
#   <div class="card">
#     <ul class="insights">
#       <li>Patient registrations show a consistent trend — monitor monthly volume for resource planning.</li>
#       <li>Top diagnoses drive most of the clinical workload — target training and drug stocking accordingly.</li>
#       <li>Referral pathway indicates the strongest downstream load (Lab vs Pharmacy vs Optical).</li>
#       <li>Low-stock alerts must be actioned promptly to avoid stockouts.</li>
#       <li>Use the forecasting module to plan drug &amp; staff demand for the coming month.</li>
#       <li>Patient clusters reveal distinct risk groups for targeted interventions.</li>
#     </ul>
#   </div>

#   <footer class="report-footer">
#     COREP Analytics Pipeline · {{ generated_at }}
#   </footer>

# </div>
# </body>
# </html>
# """


# # ============================================================
# # HTML Report
# # ============================================================
# def generate_html_report(summary, charts):
#     html = Template(HTML_TEMPLATE).render(
#         generated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
#         summary=summary,
#         charts=charts,
#     )
#     out = os.path.join(REPORT_DIR, "corep_report.html")
#     with open(out, "w", encoding="utf-8") as f:
#         f.write(html)
#     print(f"📄 Responsive HTML report → {out}")
#     return out, html


# # ============================================================
# # Word Report
# # ============================================================
# def generate_word_report(summary, chart_paths):
#     """chart_paths: list of (filename, full_path)."""
#     try:
#         from docx import Document
#         from docx.shared import Inches, Pt, RGBColor
#         from docx.enum.text import WD_ALIGN_PARAGRAPH
#     except ImportError:
#         print("⚠️  python-docx not installed. Run: pip install python-docx")
#         return None

#     doc = Document()

#     # Title
#     title = doc.add_heading("COREP Clinical Data Analysis Report", level=0)
#     title.alignment = WD_ALIGN_PARAGRAPH.CENTER

#     sub = doc.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}")
#     sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
#     sub.runs[0].font.size = Pt(10)
#     sub.runs[0].font.color.rgb = RGBColor(100, 100, 100)

#     doc.add_paragraph()

#     # Executive Summary
#     doc.add_heading("1. Executive Summary", level=1)
#     table = doc.add_table(rows=0, cols=2)
#     table.style = "Light Grid Accent 1"
#     for k, v in summary.items():
#         row = table.add_row().cells
#         row[0].text = str(k)
#         row[1].text = str(v)
#         for p in row[0].paragraphs:
#             for r in p.runs:
#                 r.bold = True

#     doc.add_page_break()

#     # Charts
#     doc.add_heading("2. Charts & Findings", level=1)
#     for filename, path in chart_paths:
#         # Chart caption
#         cap = doc.add_paragraph(pretty_name(filename))
#         cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
#         for r in cap.runs:
#             r.bold = True
#             r.font.size = Pt(11)
#             r.font.color.rgb = RGBColor(11, 74, 111)

#         # Chart image
#         try:
#             p = doc.add_paragraph()
#             p.alignment = WD_ALIGN_PARAGRAPH.CENTER
#             run = p.add_run()
#             run.add_picture(path, width=Inches(5.8))
#         except Exception as e:
#             doc.add_paragraph(f"[Chart could not be embedded: {e}]")

#         doc.add_paragraph()

#     # Key Insights
#     doc.add_page_break()
#     doc.add_heading("3. Key Insights", level=1)
#     insights = [
#         "Patient registrations show a consistent trend — monitor monthly volume for resource planning.",
#         "Top diagnoses drive most of the clinical workload — target training and drug stocking accordingly.",
#         "Referral pathway indicates the strongest downstream load (Lab vs Pharmacy vs Optical).",
#         "Low-stock alerts must be actioned promptly to avoid stockouts.",
#         "Use the forecasting module to plan drug & staff demand for the coming month.",
#         "Patient clusters reveal distinct risk groups for targeted interventions.",
#     ]
#     for item in insights:
#         doc.add_paragraph(item, style="List Bullet")

#     out = os.path.join(REPORT_DIR, "corep_report.docx")
#     doc.save(out)
#     print(f"📝 Word report → {out}")
#     return out


# # ============================================================
# # PDF (Optional, graceful fallback)
# # ============================================================
# def generate_pdf_report(html_str):
#     """Try WeasyPrint; if it fails, skip without crashing."""
#     try:
#         from weasyprint import HTML
#     except Exception as e:
#         print(f"ℹ️  PDF skipped (WeasyPrint unavailable): {e}")
#         return None

#     try:
#         out = os.path.join(REPORT_DIR, "corep_report.pdf")
#         HTML(string=html_str, base_url=os.getcwd()).write_pdf(
#             out,
#             presentational_hints=True,  # respects img width attrs [citation:12]
#         )
#         print(f"📕 PDF report → {out}")
#         return out
#     except Exception as e:
#         print(f"ℹ️  PDF skipped (WeasyPrint failed): {e}")
#         print("    → Fix: install GTK3 runtime or use conda-forge.")
#         return None


# # ============================================================
# # Main
# # ============================================================
# def generate_report(data=None):
#     """Run full report generation. If data is None, load + clean first."""
#     print("\n=== GENERATING REPORT ===")

#     if data is None:
#         import main_analysis as analysis
#         data = analysis.clean_all(analysis.load_data())

#     summary = build_summary(data)

#     # HTML uses base64 (self-contained, works offline)
#     chart_files = collect_charts()
#     charts_b64 = [(name, img_to_base64(path)) for name, path in chart_files]

#     html_path, html_str = generate_html_report(summary, charts_b64)
#     word_path = generate_word_report(summary, chart_files)
#     pdf_path = generate_pdf_report(html_str)

#     print("\n📦 Report outputs:")
#     print(f"   • HTML: {html_path}")
#     if word_path: print(f"   • Word: {word_path}")
#     if pdf_path:  print(f"   • PDF : {pdf_path}")

#     return {"html": html_path, "word": word_path, "pdf": pdf_path}


# if __name__ == "__main__":
#     generate_report()


